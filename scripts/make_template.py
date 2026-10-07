#!/usr/bin/env python3
"""
Builds data/beauty_helper_data_template.xlsx: an empty workbook with the two tables (products, reviews) and drop-down lists,
so the values that scripts/check_data.py accepts are the only ones you can pick.   Run:  python scripts/make_template.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PRODUCT_COLS = ["product_id", "product_name", "role", "brand", "size", "platform", "url", "price_sgd", "price_date",
                "skin_type_label", "official_skin_type_text", "official_benefits_text", "main_claim", "description", "ingredients"]
REVIEW_COLS = ["review_id", "product_id", "platform", "review_text", "rating", "use", "aspect", "label", "usable", "skip_reason", "source"]
WIDTHS_P = [10, 34, 11, 16, 10, 12, 36, 11, 12, 15, 30, 30, 28, 40, 40]
WIDTHS_R = [10, 11, 12, 70, 8, 8, 10, 11, 9, 24, 20]
ROWS = 400


def sheet(wb, title, cols, widths):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    return ws


def dropdown(ws, col_letter, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True, showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{ROWS}")


wb = Workbook()
wb.remove(wb.active)
guide = wb.create_sheet("how_to_fill")
for line in [
    "How to fill this workbook (details: docs/DATA_FORMAT.md)",
    "",
    "products sheet: one row per product. product_id like P01, P02 ... role = kb for a moisturizer the assistant may recommend (15 of these); train_only for any other cosmetic product you only use for classifier training.",
    "  For role = kb: price in SGD as a number with the date you noted it, skin_type_label, main_claim are required. For train_only only product_id, name, platform and role are needed.",
    "  skin_type_label: choose ONE (oily / dry / sensitive) from the official product page, not from your own feeling.",
    "  main_claim: the first and most prominent benefit on the page. official_*_text: copy the page's wording.",
    "reviews sheet: one row per review. product_id must exist in the products sheet. Copy review text only; no usernames or personal details.",
    "  use = kb    the first ~15 reviews on the product page, copied in page order, NOT selected.",
    "  use = train reviews you picked on purpose with the star filter: 1-2 star AND 4-5 star, from the SAME products.",
    "  aspect = product if the review is about the product itself (skin, texture, smell, results); service if it is only about delivery, packaging, seller, price or authenticity.",
    "  usable = no for 'not tried yet' or undecidable reviews (and for service reviews); write the reason in skip_reason.",
    "  label (positive / negative) is YOUR hand label. Required for every usable 'train' review; recommended for 'kb' reviews too.",
    "Then run:  python scripts/check_data.py --workbook data/beauty_helper_data.xlsx",
]:
    guide.append([line])
guide.column_dimensions["A"].width = 140
guide["A1"].font = Font(bold=True, size=13)

p = sheet(wb, "products", PRODUCT_COLS, WIDTHS_P)
r = sheet(wb, "reviews", REVIEW_COLS, WIDTHS_R)
dropdown(p, "C", ["kb", "train_only"])
dropdown(p, "F", ["Shopee", "Lazada", "Sephora SG"])
dropdown(p, "J", ["oily", "dry", "sensitive"])
dropdown(r, "C", ["Shopee", "Lazada", "Sephora SG"])
dropdown(r, "F", ["kb", "train"])
dropdown(r, "G", ["product", "service"])
dropdown(r, "H", ["positive", "negative"])
dropdown(r, "I", ["yes", "no"])
rating = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True, showErrorMessage=True)
r.add_data_validation(rating)
rating.add(f"E2:E{ROWS}")
for ws in (p, r):
    for row in ws.iter_rows(min_row=2, max_row=ROWS):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

out = Path(__file__).resolve().parent.parent / "data" / "beauty_helper_data_template.xlsx"
wb.save(out)
print("wrote", out)
